"""재무 내려받기 — 엑셀 파일로.

CSV는 열이 31개라 엑셀에서 열면 폭이 다 좁고 숫자가 raw로 찍혀 읽기 어렵다.
기획팀이 준 `재무실적 정리 2026 09 29.xlsx`처럼 **두 줄 머리글(분기는 묶어서)**,
천단위 구분, 퍼센트 표기, 틀 고정까지 넣어 바로 볼 수 있게 만든다.

범위는 **화장품업 전체**다. 모니터링 브랜드만 담으면 26개사뿐이라 업계 안에서
어디쯤인지 비교가 안 된다. 등록 브랜드는 '브랜드명' 칸이 채워져 구분된다.
"""

import io
import logging
from datetime import datetime

logger = logging.getLogger(__name__)

_MONEY = '#,##0'
_PCT = '0.0"%"'


def _groups(base_y: int, prev_y: int, est_y: int) -> list:
    """(그룹명, [하위라벨], 서식) — 하위라벨이 비면 단일 열."""
    return [
        ("회사명", [], None), ("브랜드명", [], None),
        (f"{base_y}년 매출", [], _MONEY),
        (f"{prev_y}년 매출", [], _MONEY),
        (f"{prev_y}년 매출 (분기별)", ["1분기", "2분기", "3분기", "4분기"], _MONEY),
        (f"{est_y}년 매출", [], _MONEY),
        (f"{est_y}년 매출 (분기별)", ["1분기", "2분기", "3분기", "4분기"], _MONEY),
        (f"{base_y}년 영업이익", [], _MONEY),
        (f"{prev_y}년 영업이익", [], _MONEY),
        (f"{prev_y}년 영업이익 (분기별)", ["1분기", "2분기", "3분기", "4분기"], _MONEY),
        (f"{est_y}년 영업이익", [], _MONEY),
        (f"{est_y}년 영업이익 (분기별)", ["1분기", "2분기", "3분기", "4분기"], _MONEY),
        (f"{base_y}년 영업이익률", [], _PCT),
        (f"{prev_y}년 영업이익률", [], _PCT),
        (f"{est_y}년 영업이익률", [], _PCT),
        (f"{base_y}년 광고비", [], _MONEY),
        (f"{prev_y}년 광고비", [], _MONEY),
        (f"{base_y}년 광고비율", [], _PCT),
        (f"{prev_y}년 광고비율", [], _PCT),
    ]


def build_workbook(rows: list, base_y: int, prev_y: int, est_y: int) -> bytes:
    """rows = get_financial_export()의 행(31칸) → .xlsx 바이트."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "재무실적"

    groups = _groups(base_y, prev_y, est_y)
    # 열 서식을 순서대로 펼친다(그룹 하나가 1칸 또는 4칸)
    fmts, widths = [], []
    for name, subs, fmt in groups:
        n = len(subs) or 1
        fmts += [fmt] * n
        if name == "회사명":
            widths += [26]
        elif name == "브랜드명":
            widths += [22]
        elif subs:
            widths += [10] * n
        else:
            widths += [13]

    ink = "1F2A44"
    head_fill = PatternFill("solid", fgColor="E8EDF5")
    sub_fill = PatternFill("solid", fgColor="F4F7FB")
    thin = Side(style="thin", color="C9D3E0")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 1행 — 무엇을 담은 파일인지. 열어본 사람이 범위·단위를 바로 알게.
    ws.cell(row=1, column=1,
            value=(f"화장품업 재무실적 · 단위 억원 · {datetime.now():%Y-%m-%d} 생성 — "
                   f"{base_y}~{prev_y}년은 NICE 확정, {est_y}년은 증권사 추정(상장사만). "
                   f"브랜드명이 있는 행이 모니터링 대상입니다."))
    ws.cell(row=1, column=1).font = Font(size=10, color="5C646E")

    # 2~3행 — 두 줄 머리글. 분기는 묶고, 단일 열은 두 줄을 합친다.
    col = 1
    for name, subs, _f in groups:
        c = ws.cell(row=2, column=col, value=name)
        c.font = Font(bold=True, size=10, color=ink)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = bd
        if subs:
            ws.merge_cells(start_row=2, start_column=col,
                           end_row=2, end_column=col + len(subs) - 1)
            for i, sub in enumerate(subs):
                sc = ws.cell(row=3, column=col + i, value=sub)
                sc.font = Font(size=9.5, color="5C646E")
                sc.fill = sub_fill
                sc.alignment = Alignment(horizontal="center")
                sc.border = bd
            col += len(subs)
        else:
            ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
            col += 1

    # 4행부터 데이터
    for ri, row in enumerate(rows, start=4):
        for ci, val in enumerate(row, start=1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.border = bd
            f = fmts[ci - 1] if ci - 1 < len(fmts) else None
            if f and isinstance(val, (int, float)):
                c.number_format = f
            if ci <= 2:
                c.alignment = Alignment(horizontal="left")
                if ci == 2 and val:
                    c.font = Font(bold=True, size=10, color=ink)
            else:
                c.alignment = Alignment(horizontal="right")

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "C4"          # 회사명·브랜드명과 머리글 고정
    ws.auto_filter.ref = f"A3:{get_column_letter(len(fmts))}{max(len(rows) + 3, 4)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
